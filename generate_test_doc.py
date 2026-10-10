"""Sinh tài liệu kiểm thử Word cho Phòng khám Đa khoa Hiếu Hải.

Chạy: python generate_test_doc.py
Kết quả: Tai-lieu-kiem-thu-he-thong.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Tai-lieu-kiem-thu-he-thong.docx"
TEAL = RGBColor(0x0F, 0x4C, 0x5C)


def shade(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=TEAL)
    return p


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run(run)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        shade(cell, "0F4C5C")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value))
            set_run(run, size=10)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


CASES = [
    ("TC-AUTH-01", "Đăng ký", "Tạo tài khoản bệnh nhân hợp lệ",
     "Chưa có tài khoản với email sắp dùng.",
     "Mở /register. Nhập họ, tên, số điện thoại từ 8 ký tự, email mới, mật khẩu từ 8 ký tự. Tick điều khoản. Bấm Tạo tài khoản.",
     "Vào cổng bệnh nhân. Tài khoản có vai trò PATIENT. Không tạo workspace hay tài khoản nhân viên."),
    ("TC-AUTH-02", "Đăng ký", "Từ chối dữ liệu thiếu hoặc mật khẩu ngắn",
     "Đang ở /register.",
     "Bỏ trống từng trường bắt buộc. Thử mật khẩu dưới 8 ký tự và số điện thoại dưới 8 ký tự.",
     "Form không gửi hoặc báo lỗi. Không tạo tài khoản."),
    ("TC-AUTH-03", "Đăng ký", "Từ chối email đã tồn tại",
     "Đã có patient@clinic.test.",
     "Đăng ký lại đúng email đó.",
     "Báo email đã được sử dụng. Không đăng nhập thành tài khoản mới."),
    ("TC-AUTH-04", "Đăng ký", "Bắt buộc đồng ý điều khoản",
     "Đang ở /register, form đã điền đủ.",
     "Không tick điều khoản rồi bấm Tạo tài khoản.",
     "Không tạo tài khoản. Có thông báo cần đồng ý điều khoản."),
    ("TC-AUTH-05", "Đăng nhập", "Bệnh nhân vào cổng đặt lịch",
     "Có patient@clinic.test / Patient123!.",
     "Mở /login, đăng nhập tài khoản bệnh nhân.",
     "URL chuyển /patient. Thấy lời chào tên bệnh nhân và nút Đặt lịch khám. Không vào được /admin."),
    ("TC-AUTH-06", "Đăng nhập", "Sai mật khẩu",
     "Đang ở /login.",
     "Nhập đúng email, sai mật khẩu.",
     "Ở lại trang đăng nhập. Báo email hoặc mật khẩu không chính xác."),
    ("TC-AUTH-07", "Đăng nhập", "Admin dùng tài khoản có sẵn",
     "Có admin@clinic.test / Admin123!. Không tự đăng ký admin.",
     "Đăng nhập tài khoản admin.",
     "Vào /admin, thấy Tổng quan, Lịch hẹn, Bác sĩ, Dịch vụ, Cơ sở."),
    ("TC-AUTH-08", "Đăng nhập", "Bệnh nhân không mở được khu quản trị",
     "Đang đăng nhập bệnh nhân.",
     "Sửa URL thành /admin.",
     "Bị đưa về /patient. Gọi API /api/admin trả 403."),
    ("TC-AUTH-09", "Đăng nhập", "Duy trì đăng nhập và đăng xuất",
     "Đăng nhập với ô Duy trì đăng nhập được chọn.",
     "Tải lại trang. Sau đó đăng xuất.",
     "Tải lại vẫn còn phiên. Sau đăng xuất về /login và không còn token."),
    ("TC-AUTH-10", "Đăng nhập", "Quên mật khẩu chỉ là thông báo trên giao diện",
     "Đang ở /login.",
     "Bấm Quên mật khẩu, nhập email, gửi.",
     "Hiện màn kiểm tra hộp thư. Không có email thật được gửi. Ghi nhận đây là giới hạn đã biết."),
    ("TC-PAT-01", "Cổng bệnh nhân", "Tìm bác sĩ theo chuyên khoa và cơ sở",
     "Đã đăng nhập bệnh nhân.",
     "Nhập chuyên khoa hoặc tên bác sĩ, chọn địa điểm, bấm Tìm kiếm. Bấm Đặt lịch trên một kết quả.",
     "Danh sách bác sĩ khớp bộ lọc. Wizard đặt lịch mở đúng bác sĩ đã chọn."),
    ("TC-PAT-02", "Đặt lịch", "Đặt lịch đủ 4 bước",
     "Đăng nhập bệnh nhân. Chọn ngày trong tuần còn giờ trống.",
     "Chọn bác sĩ và dịch vụ. Chọn ngày và một khung giờ. Nhập họ tên, số điện thoại. Xác nhận.",
     "Bước thành công hiện mã phiếu dạng MA-.... Trạng thái chờ duyệt (PENDING). Lịch xuất hiện ở Lịch hẹn của tôi."),
    ("TC-PAT-03", "Đặt lịch", "Không cho đặt khi thiếu bác sĩ, dịch vụ hoặc giờ",
     "Đang mở wizard.",
     "Bấm tiếp khi chưa chọn bác sĩ, dịch vụ, ngày hoặc khung giờ. Nhập số điện thoại dưới 8 ký tự.",
     "Dừng ở bước tương ứng và báo lỗi. Không tạo phiếu."),
    ("TC-PAT-04", "Đặt lịch", "Lịch làm việc và khung giờ",
     "Đăng nhập bệnh nhân.",
     "Chọn thứ Bảy hoặc Chủ nhật. Chọn một ngày trong tuần. So sánh giờ sáng và chiều.",
     "Cuối tuần không có giờ trống. Ngày thường có slot 30 phút trong 08:00–12:00 và 13:00–17:00. Giờ đã qua trong ngày không hiện."),
    ("TC-PAT-05", "Đặt lịch", "Không đặt trùng một khung giờ",
     "Một khung giờ của bác sĩ đã có lịch PENDING hoặc CONFIRMED.",
     "Bệnh nhân khác đặt đúng bác sĩ, ngày và giờ đó.",
     "Khung giờ không còn trong danh sách trống, hoặc API báo khung giờ đã được đặt."),
    ("TC-PAT-06", "Đặt lịch", "Hủy lịch trước 2 giờ",
     "Có lịch PENDING hoặc CONFIRMED cách giờ khám hơn 2 giờ.",
     "Ở Lịch hẹn của tôi, hủy lịch và xác nhận.",
     "Trạng thái thành Đã hủy. Có lý do hủy."),
    ("TC-PAT-07", "Đặt lịch", "Không hủy sát giờ khám",
     "Có lịch bắt đầu trong vòng chưa đầy 2 giờ, hoặc đã qua giờ.",
     "Bấm hủy lịch đó.",
     "Báo chỉ được hủy trước giờ khám tối thiểu 2 giờ. Trạng thái không đổi."),
    ("TC-PAT-08", "Đặt lịch", "Không hủy lịch đã hoàn thành hoặc đã hủy",
     "Có lịch COMPLETED hoặc CANCELLED.",
     "Thử hủy lịch đó trên giao diện hoặc API.",
     "Không chuyển được sang hủy. Thông báo lịch không còn được phép hủy."),
    ("TC-PAT-09", "Hồ sơ", "Xem và lưu hồ sơ sức khỏe",
     "Đăng nhập bệnh nhân.",
     "Mở Hồ sơ sức khỏe. Nhập cân nặng, chiều cao, huyết áp, nhịp tim, nhiệt độ, dị ứng, tiền sử. Lưu. Mở lại.",
     "Dữ liệu được lưu và hiện lại. BMI được tính khi có cân nặng và chiều cao."),
    ("TC-ADM-01", "Quản trị", "Tổng quan theo ngày và theo cơ sở",
     "Đăng nhập admin.",
     "Mở Tổng quan. Đổi bộ lọc cơ sở trên thanh trên.",
     "Có số lịch hôm nay, cần xử lý, đang hoạt động, đã khám. Đổi cơ sở thì danh sách gần đây đổi theo."),
    ("TC-ADM-02", "Duyệt lịch", "Xác nhận lịch chờ",
     "Có lịch PENDING.",
     "Mở Lịch hẹn, lọc Chờ, bấm Xác nhận.",
     "Trạng thái thành CONFIRMED. Bệnh nhân thấy lịch đã được xác nhận."),
    ("TC-ADM-03", "Duyệt lịch", "Hoàn thành lịch đã xác nhận",
     "Có lịch CONFIRMED.",
     "Bấm Hoàn thành.",
     "Trạng thái thành COMPLETED. Không còn nút xác nhận hay hoàn thành."),
    ("TC-ADM-04", "Duyệt lịch", "Admin hủy lịch đang hiệu lực",
     "Có lịch PENDING hoặc CONFIRMED.",
     "Bấm Hủy và xác nhận hộp thoại.",
     "Trạng thái CANCELLED, lý do ghi nhận hủy bởi admin."),
    ("TC-ADM-05", "Duyệt lịch", "Không đảo trạng thái đã kết thúc",
     "Có lịch COMPLETED, CANCELLED.",
     "Quan sát thao tác trên dòng đó. Nếu gọi API đổi trạng thái ngược, ghi nhận mã lỗi.",
     "Giao diện không hiện nút chuyển tiếp. API trả 422 nếu chuyển trạng thái không hợp lệ."),
    ("TC-ADM-06", "Duyệt lịch", "Lọc, tìm và xem chi tiết",
     "Có nhiều lịch khác trạng thái.",
     "Lọc từng tab. Tìm theo tên hoặc mã phiếu. Mở Chi tiết.",
     "Đúng nhóm trạng thái. Drawer hiện mã, bệnh nhân, bác sĩ, dịch vụ, ngày giờ, cơ sở."),
    ("TC-ADM-07", "Danh mục", "Thêm, sửa, xóa cơ sở",
     "Đăng nhập admin, mở Cơ sở.",
     "Thêm cơ sở với tên và địa chỉ. Sửa hotline. Xóa cơ sở vừa tạo nếu không bị ràng buộc dữ liệu.",
     "Danh sách cập nhật sau mỗi thao tác. Cơ sở mới xuất hiện ở bộ lọc và form bác sĩ."),
    ("TC-ADM-08", "Danh mục", "Thêm và sửa bác sĩ",
     "Đăng nhập admin, có ít nhất một chuyên khoa và một cơ sở.",
     "Thêm bác sĩ với họ tên, chuyên khoa, cơ sở. Sửa học vị. Lưu.",
     "Bác sĩ xuất hiện trong danh sách admin và có thể được bệnh nhân tìm thấy nếu đang hoạt động."),
    ("TC-ADM-09", "Danh mục", "Thêm và sửa dịch vụ",
     "Đăng nhập admin.",
     "Thêm dịch vụ với tên, chuyên khoa, thời lượng, giá. Sửa giá.",
     "Giá hiển thị dạng tiền Việt. Dịch vụ gắn đúng chuyên khoa khi bệnh nhân đặt lịch."),
    ("TC-ADM-10", "Phân quyền", "Chỉ admin gọi được API quản trị",
     "Có token bệnh nhân và token admin.",
     "Gọi GET /api/admin/appointments bằng từng token. Gọi khi không có token.",
     "Admin nhận 200. Bệnh nhân nhận 403. Không token nhận 401."),
]


def build():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TÀI LIỆU KIỂM THỬ HỆ THỐNG")
    set_run(run, size=20, bold=True, color=TEAL)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Phòng khám Đa khoa Hiếu Hải — Đặt lịch khám bệnh")
    set_run(run, size=14, bold=True)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Dùng cho tester thực hiện kiểm thử và viết báo cáo. Điền cột Kết quả thực tế, Đạt/Không đạt và Ghi chú.")
    set_run(run, size=11)

    add_heading(doc, "1. Mục đích và phạm vi")
    add_para(doc, "Hệ thống là website của một phòng khám. Người bệnh tự đăng ký tài khoản để đặt lịch. Admin dùng tài khoản có sẵn để duyệt lịch và quản lý bác sĩ, dịch vụ, cơ sở. Không có vai trò nhân viên, không có tạo workspace, không có đăng nhập Google Workspace.")
    add_bullets(doc, [
        "Trong phạm vi: đăng ký, đăng nhập, đặt lịch, hủy lịch, hồ sơ sức khỏe, duyệt lịch, quản lý cơ sở, bác sĩ, dịch vụ.",
        "Ngoài phạm vi: thanh toán online, gửi email quên mật khẩu thật, SMS, hồ sơ bệnh án đầy đủ của bác sĩ.",
    ])

    add_heading(doc, "2. Môi trường và tài khoản")
    add_table(
        doc,
        ["Hạng mục", "Giá trị"],
        [
            ["Giao diện local", "http://localhost:3000 (nếu cổng bận thì cổng Vite in ra terminal, thường là 3002)"],
            ["API", "http://localhost:4000/api"],
            ["Cơ sở dữ liệu", "SQLite, file expressjs/prisma/data/medical.db"],
            ["Admin có sẵn", "admin@clinic.test / Admin123!"],
            ["Bệnh nhân mẫu", "patient@clinic.test / Patient123!"],
            ["Phiên đăng nhập", "Token 8 giờ. Ô Duy trì đăng nhập giữ phiên sau khi tải lại trang."],
        ],
        [4.5, 12.5],
    )
    add_para(doc, "Chạy backend trong thư mục expressjs: npm install, npx prisma generate, npx prisma db push, npm run dev. Chạy frontend trong thư mục reactjs: npm install, npm run dev.")

    add_heading(doc, "3. Vai trò")
    add_table(
        doc,
        ["Vai trò", "Cách có tài khoản", "Việc được làm"],
        [
            ["Bệnh nhân", "Tự đăng ký trên /register", "Tìm bác sĩ, đặt lịch, xem và hủy lịch của mình, cập nhật hồ sơ sức khỏe."],
            ["Admin", "Tài khoản có sẵn, không đăng ký trên web", "Xem tổng quan, xác nhận hoặc hủy lịch, quản lý cơ sở, bác sĩ, dịch vụ."],
        ],
        [3.2, 5.5, 8.3],
    )

    add_heading(doc, "4. Chức năng cần kiểm")
    add_heading(doc, "4.1. Đăng ký và đăng nhập", 2)
    add_bullets(doc, [
        "Đăng ký gồm họ, tên, số điện thoại, email, mật khẩu và điều khoản. Mật khẩu tối thiểu 8 ký tự, số điện thoại tối thiểu 8 ký tự.",
        "Sau đăng ký vào thẳng cổng bệnh nhân.",
        "Đăng nhập bệnh nhân vào /patient. Đăng nhập admin vào /admin.",
        "Quên mật khẩu chỉ hiện thông báo trên màn hình, chưa gửi email.",
    ])
    add_heading(doc, "4.2. Cổng bệnh nhân", 2)
    add_bullets(doc, [
        "Tìm bác sĩ theo chuyên khoa hoặc tên, lọc theo cơ sở và ngày.",
        "Đặt lịch 4 bước: bác sĩ và dịch vụ, ngày và khung giờ, thông tin người khám, mã phiếu thành công.",
        "Lịch mới ở trạng thái chờ duyệt. Bệnh nhân xem lịch của mình và hủy nếu còn sớm hơn giờ khám ít nhất 2 giờ.",
        "Hồ sơ sức khỏe: cân nặng, chiều cao, huyết áp, nhịp tim, nhiệt độ, dị ứng, tiền sử, ghi chú.",
        "Bốn cơ sở mẫu: Quận 5, Bình Thạnh, TP. Thủ Đức, Cầu Giấy.",
    ])
    add_heading(doc, "4.3. Quản trị", 2)
    add_bullets(doc, [
        "Tổng quan: số lịch hôm nay, cần xử lý, đang hoạt động, đã khám, lọc theo cơ sở.",
        "Lịch hẹn: lọc trạng thái, tìm theo bệnh nhân hoặc mã phiếu, xem chi tiết.",
        "Duyệt: PENDING sang CONFIRMED hoặc CANCELLED. CONFIRMED sang COMPLETED hoặc CANCELLED.",
        "Cơ sở, bác sĩ, dịch vụ: xem, thêm, sửa. Xóa khi dữ liệu không còn bị lịch hẹn ràng buộc.",
    ])

    add_heading(doc, "5. Quy tắc nghiệp vụ")
    add_table(
        doc,
        ["Quy tắc", "Kết quả đúng"],
        [
            ["Giờ làm việc bác sĩ", "Thứ 2 đến thứ 6, 08:00–12:00 và 13:00–17:00, mỗi slot 30 phút."],
            ["Cuối tuần", "Không có khung giờ trống."],
            ["Giờ đã qua", "Không hiện trong danh sách giờ còn trống."],
            ["Trùng giờ", "Một bác sĩ không có hai lịch PENDING hoặc CONFIRMED cùng ngày và cùng giờ bắt đầu."],
            ["Hủy bởi bệnh nhân", "Chỉ hủy lịch PENDING hoặc CONFIRMED, và phải trước giờ khám ít nhất 2 giờ."],
            ["Luồng trạng thái", "PENDING → CONFIRMED hoặc CANCELLED. CONFIRMED → COMPLETED hoặc CANCELLED. Trạng thái đã kết thúc không đổi tiếp."],
            ["Mã phiếu", "Sinh tự động, dạng MA- và 8 ký tự hex."],
        ],
        [5, 12],
    )

    add_heading(doc, "6. Ca kiểm thử")
    add_para(doc, "Với mỗi ca, thực hiện đúng các bước rồi điền ba cột cuối. Đạt khi kết quả thực tế khớp kết quả mong đợi. Không đạt khi lệch. Bị chặn khi thiếu dữ liệu hoặc môi trường không chạy.")
    add_table(
        doc,
        ["Mã", "Nhóm", "Mục tiêu", "Điều kiện", "Các bước", "Kết quả mong đợi", "Kết quả thực tế", "Đạt/Không đạt", "Ghi chú"],
        [row + ("", "", "") for row in CASES],
        [2.2, 2.2, 3.2, 3.2, 4.2, 4.2, 3, 2.2, 2.4],
    )

    add_heading(doc, "7. Mẫu báo cáo")
    add_table(
        doc,
        ["Mục", "Nội dung tester điền"],
        [
            ["Người kiểm thử", ""],
            ["Ngày kiểm thử", ""],
            ["Phiên bản / commit", ""],
            ["Môi trường", "Local hoặc URL đã deploy, trình duyệt"],
            ["Tổng số ca", str(len(CASES))],
            ["Đạt", ""],
            ["Không đạt", ""],
            ["Bị chặn", ""],
            ["Lỗi nghiêm trọng", "Liệt kê mã ca và cách tái hiện"],
            ["Kết luận", "Đạt để bàn giao / chưa đạt, cần sửa"],
        ],
        [5, 12],
    )
    add_para(doc, "Với mỗi ca Không đạt, báo cáo cần có: mã ca, các bước đã làm, kết quả thực tế, ảnh màn hình, và request/response nếu lỗi nằm ở API.")

    add_heading(doc, "8. Giới hạn đã biết")
    add_bullets(doc, [
        "Quên mật khẩu không gửi email. Tester ghi nhận là giới hạn, không tính là lỗi hồi quy nếu thông báo hiện đúng.",
        "Nút thêm, sửa, xóa bác sĩ và dịch vụ trên bản cũ không làm gì. Bản hiện tại đã có form. Nếu nút không mở form thì ghi Không đạt.",
        "Tra cứu lịch bằng mã phiếu và số điện thoại có API POST /api/appointments/lookup nhưng chưa có màn hình riêng.",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
